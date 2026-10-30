from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font


def make_workbook(profile, drafts: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = profile.worksheet_name
    mapping = profile.column_mapping
    for field, column in mapping.items():
        sheet.cell(1, int(column), field.replace("_", " ").title()).font = Font(bold=True)
    for row_index, draft in enumerate(drafts, start=profile.start_row):
        for field, column in mapping.items():
            sheet.cell(row_index, int(column), draft.get(field))
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
